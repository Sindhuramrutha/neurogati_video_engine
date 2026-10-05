"""Editing a source file: taking slides or pages out of the file itself.

The only edit to a source file anywhere in the pipeline, so it is narrow on
purpose: slides of a deck and pages of a PDF can be removed, nothing can be
added or changed. The result is still the author's own file, minus what they
chose to drop.

The file is rewritten, not copied: the file in `sources/` is the one every stage
reads, and a second, edited copy beside it would leave two files claiming to be
the source. Extraction notices the new size and flags the file to be extracted
again.
"""

from __future__ import annotations

import os
from pathlib import Path
from typing import Any

from pipeline.core import OperationFailed, require_project
from pipeline.extraction_run import source_file_path


def _without_slides(path: Path, drop: list[int]) -> Any:
    """The deck with these slides (0-based) taken out, ready to save."""
    from pptx import Presentation

    presentation = Presentation(str(path))
    # A slide is listed in the presentation and reached through a relationship.
    # Unlisting it and dropping that relationship leaves its part -- and the
    # images and notes only it used -- unreachable, and python-pptx writes only
    # what is reachable, so they leave the file too.
    listed = presentation.slides._sldIdLst
    entries = list(listed)
    for index in reversed(drop):
        listed.remove(entries[index])
        presentation.part.drop_rel(entries[index].rId)
    return presentation


def _without_pages(path: Path, drop: list[int]) -> Any:
    """The PDF with these pages (0-based) taken out, ready to write."""
    from pypdf import PdfReader, PdfWriter

    # Cloned rather than rebuilt page by page, so the document's outline and
    # metadata survive; links into a removed page are cleaned away with it.
    writer = PdfWriter(clone_from=PdfReader(str(path)))
    for index in reversed(drop):
        writer.remove_page(index, clean=True)
    return writer


def _count(path: Path) -> int:
    if path.suffix.lower() == ".pptx":
        from pptx import Presentation

        return len(Presentation(str(path)).slides)
    from pypdf import PdfReader

    return len(PdfReader(str(path)).pages)


def delete_pages(slug: str, file_name: str, numbers: list[int]) -> dict[str, Any]:
    """Remove these slides or pages (1-based) from a source file, saved in place.

    Refused if it would leave nothing, or names a page the file does not have --
    a stale page asking to delete page 40 of what is now a 30-page file must not
    delete something else.
    """
    project = require_project(slug)
    path = source_file_path(project, file_name)
    suffix = path.suffix.lower()
    if suffix not in (".pptx", ".pdf"):
        raise OperationFailed(
            f"Only PowerPoint decks and PDFs can have pages removed; {path.name} is neither."
        )
    noun = "slide" if suffix == ".pptx" else "page"

    total = _count(path)
    wanted = sorted(set(numbers))
    if not wanted:
        raise OperationFailed(f"No {noun}s were chosen to delete.")
    unknown = [number for number in wanted if not 1 <= number <= total]
    if unknown:
        raise OperationFailed(
            f"{path.name} has {total} {noun}s; there is no {noun} "
            f"{', '.join(map(str, unknown))}."
        )
    if len(wanted) >= total:
        raise OperationFailed(
            f"A file needs at least one {noun}; keep one or delete the file instead."
        )

    drop = [number - 1 for number in wanted]
    edited = _without_slides(path, drop) if suffix == ".pptx" else _without_pages(path, drop)

    # Written beside it and swapped in, so a save that fails part-way leaves the
    # original untouched rather than half-written.
    saving = path.with_name(f".{path.name}.saving")
    try:
        if suffix == ".pptx":
            edited.save(str(saving))
        else:
            with saving.open("wb") as handle:
                edited.write(handle)
        os.replace(saving, path)
    finally:
        saving.unlink(missing_ok=True)

    return {
        "file_name": path.name,
        "noun": noun,
        "deleted": len(wanted),
        "remaining": total - len(wanted),
        "bytes": path.stat().st_size,
    }
