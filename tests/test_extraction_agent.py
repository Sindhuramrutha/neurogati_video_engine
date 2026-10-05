"""The Extraction agent: proposals, applying them, and undoing them."""

import io

import pytest
from PIL import Image

from pipeline import operations
from pipeline.extraction import (
    BlockKind,
    ContentBlock,
    ExtractedDocument,
    Extraction,
    ExtractionCollectionStore,
)
from pipeline.extraction_agent import (
    AgentProposal,
    AgentThreads,
    AgentThreadStore,
    AgentTurn,
    apply_proposal,
    ask_extraction_agent,
    delete_conversation,
    discard_proposal,
    read_thread,
    start_conversation,
    undo_agent_changes,
)
from pipeline.extraction_edits import BlockChange


@pytest.fixture
def project_root(tmp_path, monkeypatch):
    from pipeline.config import get_settings

    monkeypatch.setenv("PROJECTS_DIR", str(tmp_path / "projects"))
    get_settings.cache_clear()
    yield tmp_path / "projects"
    get_settings.cache_clear()


def png() -> bytes:
    buffer = io.BytesIO()
    Image.new("RGB", (40, 30), "blue").save(buffer, format="PNG")
    return buffer.getvalue()


def a_deck_collection(name="Agent"):
    """A collection with one two-slide deck: a heading, three texts and a figure."""
    slug = operations.create_project(name)["slug"]
    project = operations.require_project(slug)
    assets = project.root / "extraction" / "assets"
    assets.mkdir(parents=True, exist_ok=True)
    (assets / "fig.png").write_bytes(png())
    blocks = [
        ContentBlock(id="h1", kind=BlockKind.HEADING, text="Maps", slide_number=1),
        ContentBlock(id="t1", text="Late 1800s", slide_number=1),
        ContentBlock(id="t2", text="Animal experiments", slide_number=1),
        ContentBlock(id="f1", kind=BlockKind.FIGURE, text="A map", slide_number=1,
                     asset_path="extraction/assets/fig.png"),
        ContentBlock(id="t3", text="Second slide text", slide_number=2),
    ]
    ExtractionCollectionStore(project.root).save(Extraction(
        documents=[ExtractedDocument(file_name="deck.pptx", blocks=blocks)]))
    return slug, project


def the_agent_proposes(monkeypatch, *changes, reply="Here you go."):
    """Both ways the agent is asked -- with pictures and without -- answer this."""
    seen = {}

    def fake(model, system, user, pictures=None, **_):
        seen.update(system=system, user=user, pictures=pictures)
        return AgentProposal(reply=reply, changes=list(changes))

    monkeypatch.setattr("pipeline.llm.complete_structured", fake)
    monkeypatch.setattr("pipeline.llm.complete_structured_with_images", fake)
    return seen


def blocks_of(project):
    return ExtractionCollectionStore(project.root).load().document("deck.pptx").blocks


def test_a_proposal_changes_nothing_until_it_is_applied(project_root, monkeypatch):
    slug, project = a_deck_collection()
    the_agent_proposes(monkeypatch, BlockChange(action="merge", block_ids=["t1", "t2"]))

    thread = ask_extraction_agent(slug, "deck.pptx", "merge the first two", slide_number=1)

    agent_turn = thread["turns"][-1]
    assert agent_turn["status"] == "proposed" and agent_turn["summaries"][0].startswith("Merge 2")
    assert [b.id for b in blocks_of(project)] == ["h1", "t1", "t2", "f1", "t3"]

    apply_proposal(slug, "deck.pptx", agent_turn["id"])
    blocks = blocks_of(project)
    assert [b.id for b in blocks] == ["h1", "t1", "f1", "t3"]
    assert blocks[1].text == "Late 1800s Animal experiments"


def test_the_agent_sees_the_current_slide_in_full_and_the_rest_as_an_index(
        project_root, monkeypatch):
    slug, _ = a_deck_collection()
    seen = the_agent_proposes(monkeypatch)

    ask_extraction_agent(slug, "deck.pptx", "what is here?", slide_number=1)

    assert "Current slide: 1" in seen["user"]
    assert "[t2] paragraph: Animal experiments" in seen["user"]
    assert "act on these only if asked" in seen["user"] and "[t3]" in seen["user"]


def test_a_change_that_could_never_apply_is_dropped_with_its_reason(project_root, monkeypatch):
    slug, _ = a_deck_collection()
    the_agent_proposes(
        monkeypatch,
        BlockChange(action="merge", block_ids=["t1", "f1"]),        # a figure cannot merge
        BlockChange(action="edit_text", block_ids=["nope"], text="x"),
        BlockChange(action="change_kind", block_ids=["t1"], kind="heading"),
    )

    turn = ask_extraction_agent(slug, "deck.pptx", "tidy up", slide_number=1)["turns"][-1]

    assert [c["action"] for c in turn["changes"]] == ["change_kind"]
    assert len(turn["notes"]) == 2 and "figures stay whole" in turn["notes"][0]


def test_every_kind_of_text_change_applies(project_root, monkeypatch):
    slug, project = a_deck_collection()
    the_agent_proposes(
        monkeypatch,
        BlockChange(action="edit_text", block_ids=["h1"], text="Maps, Maps Everywhere"),
        BlockChange(action="change_kind", block_ids=["t1"], kind="list_item"),
        BlockChange(action="move", block_ids=["t3"], after_block_id="h1"),
        BlockChange(action="delete", block_ids=["t2"]),
    )
    turn = ask_extraction_agent(slug, "deck.pptx", "go", slide_number=1)["turns"][-1]

    apply_proposal(slug, "deck.pptx", turn["id"])

    blocks = blocks_of(project)
    assert [b.id for b in blocks] == ["h1", "t3", "t1", "f1"]
    assert blocks[0].text == "Maps, Maps Everywhere"
    assert blocks[2].kind == "list_item"
    assert blocks[1].slide_number == 1, "a moved block belongs where it was put"


def test_undo_puts_the_blocks_back(project_root, monkeypatch):
    slug, project = a_deck_collection()
    the_agent_proposes(monkeypatch, BlockChange(action="delete", block_ids=["t1", "t2"]))
    turn = ask_extraction_agent(slug, "deck.pptx", "delete", slide_number=1)["turns"][-1]
    applied = apply_proposal(slug, "deck.pptx", turn["id"])
    assert applied["thread"]["undo_turn_id"] == turn["id"]

    undone = undo_agent_changes(slug, "deck.pptx")

    assert [b.id for b in blocks_of(project)] == ["h1", "t1", "t2", "f1", "t3"]
    assert undone["thread"]["turns"][-1]["status"] == "undone"
    with pytest.raises(operations.OperationFailed, match="nothing to undo"):
        undo_agent_changes(slug, "deck.pptx")


def test_a_proposal_lands_whole_or_not_at_all(project_root, monkeypatch):
    slug, project = a_deck_collection()
    the_agent_proposes(
        monkeypatch,
        BlockChange(action="edit_text", block_ids=["t1"], text="changed"),
        BlockChange(action="redraw_image", block_ids=["f1"], instruction="bigger labels"),
    )
    turn = ask_extraction_agent(slug, "deck.pptx", "go", slide_number=1)["turns"][-1]

    def unreachable(prompt, **_):
        from pipeline.llm import ModelCallFailed
        raise ModelCallFailed("provider down")

    monkeypatch.setattr("pipeline.llm.draw_image", unreachable)
    with pytest.raises(operations.OperationFailed, match="provider down"):
        apply_proposal(slug, "deck.pptx", turn["id"])

    assert blocks_of(project)[1].text == "Late 1800s", "the text edit did not land alone"
    assert read_thread(slug, "deck.pptx")["turns"][-1]["status"] == "proposed"


def test_a_redrawn_and_an_edited_picture_are_recorded_as_generated(project_root, monkeypatch):
    slug, project = a_deck_collection()
    sent = {}
    monkeypatch.setattr("pipeline.llm.draw_image", lambda prompt, **_: png())

    def edit(image, mime, instruction):
        sent.update(bytes=image, mime=mime, instruction=instruction)
        return png()

    monkeypatch.setattr("pipeline.llm.edit_image", edit)
    for action in ("redraw_image", "edit_image"):
        the_agent_proposes(monkeypatch, BlockChange(
            action=action, block_ids=["f1"], instruction="crop to the map"))
        turn = ask_extraction_agent(slug, "deck.pptx", "picture", slide_number=1)["turns"][-1]
        apply_proposal(slug, "deck.pptx", turn["id"])
        figure = next(b for b in blocks_of(project) if b.id == "f1")
        assert figure.asset_origin == "generated"
        assert (project.root / figure.asset_path).exists()

    assert sent["mime"] == "image/png" and sent["instruction"].startswith("crop to the map.")
    assert "Do not add a title" in sent["instruction"], "an edit keeps the rest as it is"


def test_a_discarded_proposal_cannot_be_applied(project_root, monkeypatch):
    slug, _ = a_deck_collection()
    the_agent_proposes(monkeypatch, BlockChange(action="delete", block_ids=["t1"]))
    turn = ask_extraction_agent(slug, "deck.pptx", "delete", slide_number=1)["turns"][-1]

    discard_proposal(slug, "deck.pptx", turn["id"])

    with pytest.raises(operations.OperationFailed, match="not open"):
        apply_proposal(slug, "deck.pptx", turn["id"])


def test_the_conversation_runs_over_http(project_root, monkeypatch):
    from litestar.testing import TestClient

    from pipeline.web import routes_project
    from pipeline.web.app import create_app

    slug, _ = a_deck_collection("Over HTTP")
    the_agent_proposes(monkeypatch, BlockChange(action="delete", block_ids=["t3"]))
    routes_project._service = None
    with TestClient(app=create_app()) as client:
        asked = client.post(f"/api/projects/{slug}/extraction/agent",
                            json={"file_name": "deck.pptx", "message": "drop slide 2's text",
                                  "slide_number": 2})
        turn_id = asked.json()["turns"][-1]["id"]
        applied = client.post(f"/api/projects/{slug}/extraction/agent/{turn_id}/apply",
                              json={"file_name": "deck.pptx"})
        thread = client.get(f"/api/projects/{slug}/extraction/agent",
                            params={"file_name": "deck.pptx"})
    routes_project._service = None

    assert asked.status_code == 201, asked.text
    assert applied.status_code == 201, applied.text
    assert applied.json()["extraction"]["freshness"]["pending_confirmation"] is True
    assert thread.json()["turns"][-1]["status"] == "applied"


def a_real_deck(project, name="deck.pptx"):
    from pptx import Presentation
    from pptx.util import Inches

    presentation = Presentation()
    for title in ("Maps", "Second"):
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1)).text_frame.text = title
    project.sources_dir.mkdir(parents=True, exist_ok=True)
    presentation.save(str(project.sources_dir / name))


def test_the_agent_is_shown_the_whole_slide_and_its_figures(project_root, monkeypatch):
    slug, project = a_deck_collection()
    a_real_deck(project)
    seen = the_agent_proposes(monkeypatch)

    turn = ask_extraction_agent(slug, "deck.pptx", "what is in the picture?", slide_number=1)
    labels = [label for label, _ in seen["pictures"]]

    assert labels[0] == "Picture 1: the whole of slide 1, as it looks."
    assert labels[1].startswith("Picture 2: figure [f1]")
    assert all(picture[:2] == bytes([0xFF, 0xD8]) for _, picture in seen["pictures"]), \
        "sent as JPEG"
    assert turn["turns"][-1]["looked_at"] == "slide 1, 1 figure"

    # Drawn once, then kept: asking again does not draw the slide again.
    kept = list((project.root / "extraction" / "agent-slides").glob("*.jpg"))
    ask_extraction_agent(slug, "deck.pptx", "and now?", slide_number=1)
    assert list((project.root / "extraction" / "agent-slides").glob("*.jpg")) == kept


def test_without_the_vision_model_the_agent_answers_from_the_text(project_root, monkeypatch):
    from pipeline.llm import ModelCallFailed

    slug, _ = a_deck_collection()
    the_agent_proposes(monkeypatch, reply="From the text: a map.")

    def cannot_see(*_args, **_kwargs):
        raise ModelCallFailed("images are not supported by this model")

    monkeypatch.setattr("pipeline.llm.complete_structured_with_images", cannot_see)
    turn = ask_extraction_agent(slug, "deck.pptx", "what is it?", slide_number=1)["turns"][-1]

    assert turn["text"] == "From the text: a map."
    assert turn["looked_at"] == ""
    assert "could not look at the pictures" in turn["notes"][0]


def test_a_new_chat_starts_empty_and_the_old_one_stays_in_history(project_root, monkeypatch):
    slug, _ = a_deck_collection()
    the_agent_proposes(monkeypatch)
    first = ask_extraction_agent(slug, "deck.pptx", "first question", slide_number=1)

    fresh = start_conversation(slug, "deck.pptx")
    assert fresh["turns"] == [] and fresh["conversation_id"] != first["conversation_id"]
    reused = start_conversation(slug, "deck.pptx")["conversation_id"]
    assert reused == fresh["conversation_id"], "an empty chat is reused rather than piling up"

    ask_extraction_agent(slug, "deck.pptx", "second question", slide_number=1,
                         conversation_id=fresh["conversation_id"])
    history = read_thread(slug, "deck.pptx")["conversations"]
    assert [c["title"] for c in history] == ["second question", "first question"]

    reopened = read_thread(slug, "deck.pptx", first["conversation_id"])
    assert reopened["turns"][0]["text"] == "first question"


def test_a_chat_from_before_conversations_is_kept_as_the_first_one(project_root):
    slug, project = a_deck_collection()
    AgentThreadStore(project.root).save(AgentThreads(threads={"deck.pptx": [
        AgentTurn(role="person", text="hi", at="2026-09-01T00:00:00+00:00"),
        AgentTurn(role="agent", text="hello", at="2026-09-01T00:00:01+00:00"),
    ]}))

    thread = read_thread(slug, "deck.pptx")

    assert [t["text"] for t in thread["turns"]] == ["hi", "hello"]
    assert thread["conversations"][0]["title"] == "hi"


def test_a_chat_can_be_deleted_from_the_history(project_root, monkeypatch):
    slug, project = a_deck_collection()
    the_agent_proposes(monkeypatch, BlockChange(action="delete", block_ids=["t3"]))
    first = ask_extraction_agent(slug, "deck.pptx", "drop it", slide_number=2)
    apply_proposal(slug, "deck.pptx", first["turns"][-1]["id"])
    second = start_conversation(slug, "deck.pptx")
    ask_extraction_agent(slug, "deck.pptx", "keep this one", slide_number=1,
                         conversation_id=second["conversation_id"])

    left = delete_conversation(slug, "deck.pptx", first["conversation_id"])

    assert [c["title"] for c in left["conversations"]] == ["keep this one"]
    assert "t3" not in [b.id for b in blocks_of(project)], "what it applied stays applied"
    with pytest.raises(operations.OperationFailed, match="not in the history"):
        delete_conversation(slug, "deck.pptx", first["conversation_id"])


def test_a_new_picture_is_added_after_the_block_it_names(project_root, monkeypatch):
    slug, project = a_deck_collection()
    prompts = []
    monkeypatch.setattr("pipeline.llm.draw_image",
                        lambda prompt, **_: prompts.append(prompt) or png())
    the_agent_proposes(monkeypatch, BlockChange(
        action="add_image", after_block_id="h1", text="The visual cortex",
        instruction="a brain seen from the side, the visual cortex highlighted"))
    turn = ask_extraction_agent(slug, "deck.pptx", "add an AI image", slide_number=1)["turns"][-1]
    assert turn["summaries"][0].startswith("Add a new picture after “Maps”")

    apply_proposal(slug, "deck.pptx", turn["id"])

    blocks = blocks_of(project)
    added = blocks[1]
    assert [b.id for b in blocks][0] == "h1" and added.kind == "figure"
    assert added.text == "The visual cortex" and added.asset_origin == "generated"
    assert added.slide_number == 1 and (project.root / added.asset_path).exists()
    assert "the visual cortex highlighted" in prompts[0]


def test_a_new_picture_needs_a_place_and_a_description(project_root, monkeypatch):
    slug, _ = a_deck_collection()
    the_agent_proposes(monkeypatch, BlockChange(action="add_image", after_block_id="h1"))

    turn = ask_extraction_agent(slug, "deck.pptx", "add one", slide_number=1)["turns"][-1]

    assert turn["changes"] == [] and "description of what to draw" in turn["notes"][0]


def test_pictures_are_read_from_the_chat_reply_first(monkeypatch):
    """Gemini, as TokenRouter serves it, answers "model not found" on the image
    endpoint and puts the picture in a chat reply instead."""
    import base64
    from types import SimpleNamespace

    from pipeline import llm

    encoded = base64.b64encode(png()).decode()
    calls = []

    class Chat:
        def create(self, **kwargs):
            calls.append(kwargs)
            message = SimpleNamespace(content=f"![image](data:image/png;base64,{encoded})",
                                      model_extra={})
            return SimpleNamespace(choices=[SimpleNamespace(message=message)])

    class Images:
        def generate(self, **_):
            raise AssertionError("the image endpoint is not needed when the reply has the picture")

    client = SimpleNamespace(chat=SimpleNamespace(completions=Chat()), images=Images())
    monkeypatch.setattr(llm, "_build_client", lambda *_: client)

    assert llm.draw_image("a red circle") == png()
    assert llm.edit_image(png(), "image/png", "make it blue") == png()
    assert calls[0]["extra_body"] == {"modalities": ["image", "text"]}
    assert calls[1]["messages"][0]["content"][1]["type"] == "image_url"


# --- the slide-by-slide review ----------------------------------------------------


def test_a_slide_review_goes_in_its_own_conversation_one_answer_per_slide(
        project_root, monkeypatch):
    from pipeline.extraction_slide_review import review_slide

    slug, project = a_deck_collection()
    ask_extraction_agent(slug, "deck.pptx", "hello", slide_number=1)
    seen = the_agent_proposes(monkeypatch, reply="Slide 2 matches.")

    review = review_slide(slug, "deck.pptx", slide_number=2)
    assert review["kind"] == "review" and len(review["conversations"]) == 2
    assert "Current slide: 2" in seen["user"] and "Review what was extracted" in seen["user"]

    the_agent_proposes(monkeypatch, BlockChange(action="merge", block_ids=["t1", "t2"]))
    review = review_slide(slug, "deck.pptx", slide_number=1,
                          conversation_id=review["conversation_id"])
    assert [t["slide_number"] for t in review["turns"]] == [1, 2], "kept in slide order"
    assert all(t["role"] == "agent" for t in review["turns"])
    assert review["turns"][0]["status"] == "proposed"

    # Reviewed again, slide 1's open answer is replaced, not repeated.
    the_agent_proposes(monkeypatch, reply="Slide 1 matches.")
    review = review_slide(slug, "deck.pptx", slide_number=1,
                          conversation_id=review["conversation_id"])
    assert [t["text"] for t in review["turns"]] == ["Slide 1 matches.", "Slide 2 matches."]
    assert [b.id for b in blocks_of(project)] == ["h1", "t1", "t2", "f1", "t3"], \
        "a review only proposes"


def test_an_applied_review_answer_is_kept_when_the_slide_is_reviewed_again(
        project_root, monkeypatch):
    from pipeline.extraction_slide_review import review_slide

    slug, _ = a_deck_collection()
    the_agent_proposes(monkeypatch, BlockChange(action="merge", block_ids=["t1", "t2"]))
    review = review_slide(slug, "deck.pptx", slide_number=1)
    apply_proposal(slug, "deck.pptx", review["turns"][0]["id"])

    the_agent_proposes(monkeypatch, reply="Now it matches.")
    review = review_slide(slug, "deck.pptx", slide_number=1,
                          conversation_id=review["conversation_id"])
    assert [t["status"] for t in review["turns"]] == ["applied", ""]


# --- the picture procedure --------------------------------------------------------


def a_deck_with_a_placed_picture(name="Placed"):
    """The deck collection, with slide 1's layout showing the figure's picture."""
    from pipeline.extraction import LayoutShape, SlideLayout

    slug, project = a_deck_collection(name)
    store = ExtractionCollectionStore(project.root)
    collection = store.load()
    collection.documents[0].slides = [SlideLayout(number=1, shapes=[
        LayoutShape(kind="title", text="Maps", left=0.1, top=0.05, width=0.8, height=0.1),
        LayoutShape(kind="picture", asset_path="extraction/assets/fig.png",
                    left=0.7, top=0.1, width=0.2, height=0.8),
    ])]
    collection.documents[0].blocks.append(ContentBlock(
        id="n1", text="[speaker notes] A young man at a cocktail party", slide_number=1))
    store.save(collection)
    return slug, project


def red_png() -> bytes:
    buffer = io.BytesIO()
    Image.new("RGB", (40, 30), "red").save(buffer, format="PNG")
    return buffer.getvalue()


def test_a_redrawn_picture_takes_the_old_ones_place_on_the_slide_and_undo_puts_it_back(
        project_root, monkeypatch):
    slug, project = a_deck_with_a_placed_picture()
    monkeypatch.setattr("pipeline.llm.draw_image", lambda prompt, **_: red_png())
    the_agent_proposes(monkeypatch, BlockChange(
        action="redraw_image", block_ids=["f1"], instruction="a clean map"))
    turn = ask_extraction_agent(slug, "deck.pptx", "redraw it", slide_number=1)["turns"][-1]

    apply_proposal(slug, "deck.pptx", turn["id"])
    document = ExtractionCollectionStore(project.root).load().document("deck.pptx")
    figure = next(b for b in document.blocks if b.id == "f1")
    assert figure.asset_path != "extraction/assets/fig.png"
    slot = document.slides[0].shapes[1]
    assert slot.asset_path == figure.asset_path, "the layout shows the new picture, not the old one"

    undo_agent_changes(slug, "deck.pptx")
    document = ExtractionCollectionStore(project.root).load().document("deck.pptx")
    assert document.slides[0].shapes[1].asset_path == "extraction/assets/fig.png"


def test_an_uploaded_picture_takes_the_old_ones_place_on_the_slide(project_root):
    from pipeline.extraction_review import replace_figure

    slug, project = a_deck_with_a_placed_picture("Uploaded")
    replace_figure(slug, "f1", red_png(), ".png")

    document = ExtractionCollectionStore(project.root).load().document("deck.pptx")
    figure = next(b for b in document.blocks if b.id == "f1")
    assert figure.asset_origin == "uploaded"
    assert document.slides[0].shapes[1].asset_path == figure.asset_path


def test_a_picture_prompt_keeps_the_slide_text_out_of_the_picture(project_root):
    """The slide's text was once drawn into a picture -- title bar, paragraph and
    a speaker-notes quote -- because it was sent with nothing saying why."""
    from pipeline.figure_pictures import NO_SLIDE_TEXT, picture_prompt

    _slug, project = a_deck_with_a_placed_picture("Prompt")
    document = ExtractionCollectionStore(project.root).load().document("deck.pptx")
    figure = next(b for b in document.blocks if b.id == "f1")

    prompt = picture_prompt(document, figure, "a map of the cortex", "highlight V1")

    assert "For your understanding only" in prompt and NO_SLIDE_TEXT in prompt
    assert "Animal experiments" in prompt, "the slide's text is there to understand"
    assert "cocktail party" not in prompt, "speaker notes are not what the slide shows"
    assert "Second slide text" not in prompt, "only the picture's own slide"
    assert "Shape: portrait" in prompt, "the slot is taller than wide"
