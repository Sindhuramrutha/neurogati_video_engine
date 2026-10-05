"""Removing slides from a source deck, and pages from a source PDF."""

import pytest

from pipeline import operations
from pipeline.deck_edit import delete_pages


@pytest.fixture
def project_root(tmp_path, monkeypatch):
    from pipeline.config import get_settings

    monkeypatch.setenv("PROJECTS_DIR", str(tmp_path / "projects"))
    get_settings.cache_clear()
    yield tmp_path / "projects"
    get_settings.cache_clear()


def a_deck(path, slides):
    """A .pptx whose slides say which they are, so their order can be checked."""
    from pptx import Presentation
    from pptx.util import Inches

    presentation = Presentation()
    for number in range(1, slides + 1):
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1))
        box.text_frame.text = f"Heading number {number}"
    presentation.save(str(path))


def deck_project(name, slides=5):
    slug = operations.create_project(name)["slug"]
    project = operations.require_project(slug)
    project.sources_dir.mkdir(parents=True, exist_ok=True)
    a_deck(project.sources_dir / "deck.pptx", slides=slides)
    return slug, project


def headings(path):
    from pptx import Presentation

    return [
        next(shape.text_frame.text for shape in slide.shapes if shape.has_text_frame)
        for slide in Presentation(str(path)).slides
    ]


def test_the_chosen_slides_leave_the_file_and_the_rest_keep_their_order(project_root):
    slug, project = deck_project("Trim")

    result = delete_pages(slug, "deck.pptx", [2, 4])

    assert result["deleted"] == 2 and result["remaining"] == 3 and result["noun"] == "slide"
    assert headings(project.sources_dir / "deck.pptx") == [
        "Heading number 1", "Heading number 3", "Heading number 5",
    ]
    assert not list(project.sources_dir.glob(".*.saving")), "no half-written file is left behind"


def test_extraction_sees_the_trimmed_deck_as_changed(project_root):
    slug, project = deck_project("Changed")
    operations.run_extraction(slug, "deck.pptx")

    delete_pages(slug, "deck.pptx", [1])

    assert operations.extraction_freshness(project)["changed"] == ["deck.pptx"]


def test_a_deck_cannot_lose_every_slide(project_root):
    slug, project = deck_project("Keep One", slides=2)

    with pytest.raises(operations.OperationFailed, match="at least one slide"):
        delete_pages(slug, "deck.pptx", [1, 2])
    assert len(headings(project.sources_dir / "deck.pptx")) == 2


def test_a_slide_the_deck_does_not_have_is_refused_and_nothing_changes(project_root):
    slug, project = deck_project("Stale Page", slides=3)

    with pytest.raises(operations.OperationFailed, match="no slide 7"):
        delete_pages(slug, "deck.pptx", [1, 7])
    assert len(headings(project.sources_dir / "deck.pptx")) == 3


def test_only_decks_and_pdfs_can_have_pages_removed(project_root):
    slug, project = deck_project("Not A Deck")
    (project.sources_dir / "notes.md").write_text("# Notes", encoding="utf-8")

    with pytest.raises(operations.OperationFailed, match="Only PowerPoint decks and PDFs"):
        delete_pages(slug, "notes.md", [1])


def test_pages_are_removed_over_http(project_root):
    from litestar.testing import TestClient

    from pipeline.web import routes_project
    from pipeline.web.app import create_app

    slug, project = deck_project("Over HTTP")
    routes_project._service = None
    with TestClient(app=create_app()) as client:
        response = client.post(f"/api/projects/{slug}/sources/deck.pptx/pages/delete",
                               json={"pages": [5]})
        refused = client.post(f"/api/projects/{slug}/sources/deck.pptx/pages/delete",
                              json={"pages": [1, 2, 3, 4]})
    routes_project._service = None

    assert response.status_code == 201, response.text
    assert response.json()["remaining"] == 4
    assert refused.status_code == 400


def a_pdf(path, pages):
    """A PDF whose pages have different widths, so their order can be read back."""
    from pypdf import PdfWriter

    writer = PdfWriter()
    for number in range(1, pages + 1):
        writer.add_blank_page(width=100 * number, height=100)
    with path.open("wb") as handle:
        writer.write(handle)


def page_widths(path):
    from pypdf import PdfReader

    return [round(float(page.mediabox.width)) for page in PdfReader(str(path)).pages]


def test_the_chosen_pages_leave_a_pdf_and_the_rest_keep_their_order(project_root):
    slug, project = deck_project("Trim PDF")
    a_pdf(project.sources_dir / "notes.pdf", pages=4)

    result = delete_pages(slug, "notes.pdf", [1, 3])

    assert result["noun"] == "page" and result["remaining"] == 2
    assert page_widths(project.sources_dir / "notes.pdf") == [200, 400]


def test_a_pdf_cannot_lose_every_page(project_root):
    slug, project = deck_project("Keep A Page")
    a_pdf(project.sources_dir / "notes.pdf", pages=2)

    with pytest.raises(operations.OperationFailed, match="at least one page"):
        delete_pages(slug, "notes.pdf", [1, 2])
    assert len(page_widths(project.sources_dir / "notes.pdf")) == 2
