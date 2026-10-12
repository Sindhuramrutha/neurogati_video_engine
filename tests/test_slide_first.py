"""The deck's own slides, in the deck's own order.

A designed slide put through the old path came out as the renderer's template
with one extracted figure in a box -- the deck's background, typography and
layout all discarded. And nothing guaranteed the video even followed the deck:
sections were grouped by a model, so slides could be merged, reordered or
dropped and nothing said so.

Two promises are tested here, because both were broken:

1. **Every slide appears, in the deck's order.** Scenes may outnumber slides --
   a dense slide is narrated in several beats -- but never the reverse, and the
   slide a scene belongs to never goes backwards.
2. **A slide is shown, not rebuilt.** Its own background, pictures and text, and
   nothing of the renderer's drawn over them.
"""

from __future__ import annotations

import io
from typing import Any

import pytest
from PIL import Image, ImageDraw

from pipeline import operations


@pytest.fixture
def project_root(tmp_path, monkeypatch):
    from pipeline.config import get_settings

    monkeypatch.setenv("PROJECTS_DIR", str(tmp_path))
    get_settings.cache_clear()
    yield tmp_path
    get_settings.cache_clear()


def a_deck(path, slides: int = 5, heading_size: int = 40):
    """A real .pptx with a heading in a plain text box on every slide.

    A text box, not a title placeholder, because that is how designed decks are
    actually built -- and it is the case that must not come through as
    "Slide 7".
    """
    from pptx import Presentation
    from pptx.util import Emu, Inches, Pt

    presentation = Presentation()
    presentation.slide_width = Inches(16)
    presentation.slide_height = Inches(9)
    blank = presentation.slide_layouts[6]

    for number in range(1, slides + 1):
        slide = presentation.slides.add_slide(blank)
        box = slide.shapes.add_textbox(Inches(1), Inches(0.5), Inches(12), Inches(1.2))
        run = box.text_frame.paragraphs[0].add_run()
        run.text = f"Heading number {number}"
        run.font.size = Pt(heading_size)
        run.font.bold = True

        body = slide.shapes.add_textbox(Inches(1), Inches(2.5), Inches(12), Inches(4))
        written = body.text_frame.paragraphs[0].add_run()
        written.text = (f"Slide {number} explains a particular idea in enough words "
                        f"to be worth narrating for a while.")
        written.font.size = Pt(16)
        assert Emu(1).inches >= 0    # the import is load-bearing above
    presentation.save(str(path))
    return path


def a_deck_project(name: str, slides: int = 5):
    slug = operations.create_project(name)["slug"]
    project = operations.require_project(slug)
    (project.root / "sources").mkdir(parents=True, exist_ok=True)
    a_deck(project.root / "sources" / "deck.pptx", slides=slides)
    operations.run_extraction(slug, "deck.pptx")
    operations.confirm_extraction(slug)
    return slug, project


def slides_in(extraction, file_name: str = "") -> list[dict[str, Any]]:
    """Every slide of every deck in the extraction, in order.

    Keyed on the slide number the blocks carry, not on their heading text: two
    slides titled the same are two slides, and the title is all a heading path
    records for a deck.
    """
    found: list[dict[str, Any]] = []
    for document in extraction.documents:
        if file_name and document.file_name != file_name:
            continue
        numbers = sorted({
            block.slide_number for block in document.blocks if block.slide_number
        })
        for number in numbers:
            blocks = [b for b in document.blocks if b.slide_number == number]
            heading = next((b.text for b in blocks if b.kind == "heading"), "")
            words = [b.text for b in blocks if b.kind != "heading" and b.text.strip()]
            found.append({
                "file_name": document.file_name,
                "number": number,
                "title": heading or f"Slide {number}",
                "text": "\n".join(words),
                "figure_count": sum(1 for b in blocks if b.kind == "figure"),
            })
    return found


def plan_from_slides(slug: str) -> dict:
    """One section per slide, in the deck's order: a plan to test later layers
    against, without the planning agent."""
    from pipeline.extraction import ExtractionStore
    from pipeline.lesson_plan import save_plan
    from pipeline.planning import Plan, PlanScene, PlanSection, SourceReference

    project = operations.require_project(slug)
    slides = slides_in(ExtractionStore(project.root).load())
    save_plan(project, Plan(sections=[PlanSection(
        id=f"slide-{slide['number']:03d}", title=slide["title"], summary=slide["title"],
        scenes=[PlanScene(does=slide["title"], slides=[slide["number"]], seconds=20)],
        sources=[SourceReference(file_name=slide["file_name"], heading_path=[slide["title"]],
                                 slide_number=slide["number"])])
        for slide in slides]))
    return operations.read_artifact(slug, "plan")


# --- every slide, in order ----------------------------------------------------------


def test_the_deck_s_own_heading_survives_even_without_a_title_placeholder(project_root):
    """Designed decks set headings as plain text boxes. No slide comes
    through as "Slide 7", which would lose the one line the author most wanted kept."""
    slug, _project = a_deck_project("Headings", slides=3)

    from pipeline.extraction import ExtractionStore

    project = operations.require_project(slug)
    slides = slides_in(ExtractionStore(project.root).load())

    assert [s["title"] for s in slides] == [f"Heading number {n}" for n in (1, 2, 3)]
    assert not any(s["title"].startswith("Slide ") for s in slides)


def test_a_kicker_above_a_title_does_not_become_the_heading(project_root):
    """"4.Chapter" over "Networks that Learn" is a label above a title. The
    larger of two near-equal sizes is not reliably the heading; the longer is."""
    from pptx import Presentation
    from pptx.util import Inches, Pt

    from pipeline.extraction import ExtractionStore

    slug = operations.create_project("Kicker")["slug"]
    project = operations.require_project(slug)
    (project.root / "sources").mkdir(parents=True, exist_ok=True)

    presentation = Presentation()
    presentation.slide_width, presentation.slide_height = Inches(16), Inches(9)
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    for text, size, top in (("4.Chapter", 40, 0.5), ("Networks that Learn", 39, 1.3)):
        box = slide.shapes.add_textbox(Inches(7), Inches(top), Inches(6), Inches(0.8))
        run = box.text_frame.paragraphs[0].add_run()
        run.text, run.font.size = text, Pt(size)
    presentation.save(str(project.root / "sources" / "deck.pptx"))
    operations.run_extraction(slug, "deck.pptx")
    operations.confirm_extraction(slug)

    titles = [s["title"] for s in slides_in(ExtractionStore(project.root).load())]
    assert titles[0] == "Networks that Learn"


# --- the slide is shown, not rebuilt ------------------------------------------------


def test_a_composed_slide_carries_the_deck_s_own_pixels(project_root):
    """Not "was handed a path" -- drawn. A recognisable mark on the slide has to
    survive into the frame."""
    from pptx import Presentation
    from pptx.util import Inches

    from pipeline.slide_frame import compose_slide

    slug = operations.create_project("Pixels")["slug"]
    project = operations.require_project(slug)
    (project.root / "sources").mkdir(parents=True, exist_ok=True)

    picture = Image.new("RGB", (800, 600), (0, 0, 255))
    ImageDraw.Draw(picture).rectangle([(10, 10), (790, 590)], outline=(0, 0, 200), width=8)
    buffer = io.BytesIO()
    picture.save(buffer, format="PNG")

    presentation = Presentation()
    presentation.slide_width, presentation.slide_height = Inches(16), Inches(9)
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    buffer.seek(0)
    slide.shapes.add_picture(buffer, Inches(0), Inches(0), Inches(16), Inches(9))
    deck = project.root / "sources" / "deck.pptx"
    presentation.save(str(deck))

    composed = compose_slide(deck, 1, project.root / "slide.png")
    with Image.open(composed.path) as frame:
        blues = sum(1 for pixel in frame.convert("RGB").getdata()
                    if pixel[2] > 150 > pixel[0])

    assert frame.size == (1920, 1080)
    assert blues > 1_500_000, f"the slide's own artwork is not on the frame ({blues} px)"


