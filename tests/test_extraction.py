"""Reading source documents: order, figures, and knowing when it is out of date.

Each test here is one of three defects the Extraction agent reported against a
real project -- a table filed under the wrong section, a figure that was never
extracted, and a deck that was uploaded after the last run and silently left out.
The documents are built in the test so every assertion is about a structure we
know, not about whatever a sample file happens to contain.
"""

import io
import json

import docx
import pytest
from PIL import Image
from pptx import Presentation
from pptx.util import Inches

from pipeline import operations, orchestrator
from pipeline.extraction import (
    AssetSink,
    BlockKind,
    ContentBlock,
    ExtractedDocument,
    Extraction,
    read_docx,
    read_pptx,
)


def png_bytes(colour: str = "red", size: int = 24) -> bytes:
    buffer = io.BytesIO()
    Image.new("RGB", (size, size), colour).save(buffer, format="PNG")
    return buffer.getvalue()


@pytest.fixture
def project_root(tmp_path, monkeypatch):
    from pipeline.config import get_settings

    monkeypatch.setenv("PROJECTS_DIR", str(tmp_path / "projects"))
    get_settings.cache_clear()
    yield tmp_path / "projects"
    get_settings.cache_clear()


def chapter_with_a_table_and_a_figure(path, figure_colour="red"):
    """Two sections; the table and the figure both belong to the FIRST one."""
    document = docx.Document()
    document.add_heading("2.2 Electrodes", level=1)
    document.add_heading("2.2.9 Comparison of sensors", level=2)
    document.add_paragraph("The table below compares them.")

    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Sensor"
    table.cell(0, 1).text = "Noise"
    table.cell(1, 0).text = "Wet"
    table.cell(1, 1).text = "Low"

    picture = path.parent / f"figure-{figure_colour}.png"
    picture.write_bytes(png_bytes(figure_colour))
    document.add_picture(str(picture), width=docx.shared.Inches(1))
    document.add_paragraph("Figure 2.1 The acquisition chain.")

    document.add_heading("2.3 Summary", level=1)
    document.add_paragraph("In closing.")
    document.save(path)
    return path


# --- document order ------------------------------------------------------------


def test_a_table_stays_where_it_is_in_the_document(tmp_path):
    """python-docx keeps paragraphs and tables in two separate lists. Reading
    one then the other put every table at the end of the document."""
    blocks = read_docx(chapter_with_a_table_and_a_figure(tmp_path / "chapter.docx"))
    kinds = [block.kind for block in blocks]

    table_at = kinds.index(BlockKind.TABLE)
    summary_at = next(
        index for index, block in enumerate(blocks) if block.text == "2.3 Summary"
    )
    assert table_at < summary_at, "the table was moved after a later section"


def test_a_table_is_filed_under_its_own_section(tmp_path):
    """The reported defect: the comparison table from 2.2.9 carried the heading
    path of the closing summary, so any claim quoting it pointed at the wrong
    section."""
    blocks = read_docx(chapter_with_a_table_and_a_figure(tmp_path / "chapter.docx"))
    table = next(block for block in blocks if block.kind == BlockKind.TABLE)

    assert table.heading_path == ["2.2 Electrodes", "2.2.9 Comparison of sensors"]
    assert "Wet | Low" in table.text


# --- figures -------------------------------------------------------------------


def test_a_word_figure_is_extracted_as_an_image(tmp_path):
    sink = AssetSink(tmp_path / "assets")
    blocks = read_docx(chapter_with_a_table_and_a_figure(tmp_path / "chapter.docx"), sink)

    figures = [block for block in blocks if block.kind == BlockKind.FIGURE]
    assert len(figures) == 1
    assert figures[0].heading_path[-1] == "2.2.9 Comparison of sensors"

    stored = tmp_path / "assets" / figures[0].asset_path.split("/")[-1]
    assert stored.exists()
    assert stored.read_bytes()[:8] == b"\x89PNG\r\n\x1a\n"


def test_reading_without_a_sink_writes_nothing(tmp_path):
    """Figures are opt-in, so a reader stays a pure function of its file."""
    before = set(tmp_path.rglob("*"))
    blocks = read_docx(chapter_with_a_table_and_a_figure(tmp_path / "chapter.docx"))
    written = set(tmp_path.rglob("*")) - before - {tmp_path / "chapter.docx",
                                                   tmp_path / "figure-red.png"}

    assert written == set()
    figure = next(block for block in blocks if block.kind == BlockKind.FIGURE)
    assert figure.asset_path == ""


def test_a_slide_picture_is_extracted_as_a_figure(tmp_path):
    picture = tmp_path / "diagram.png"
    picture.write_bytes(png_bytes("blue"))

    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[5])
    slide.shapes.title.text = "Neural signal chain"
    slide.shapes.add_picture(str(picture), Inches(1), Inches(2), width=Inches(2))
    deck.save(tmp_path / "deck.pptx")

    sink = AssetSink(tmp_path / "assets")
    blocks = read_pptx(tmp_path / "deck.pptx", sink)
    figure = next(block for block in blocks if block.kind == BlockKind.FIGURE)

    assert figure.heading_path == ["Neural signal chain"]
    assert "slide 1" in figure.text
    assert (tmp_path / "assets" / figure.asset_path.split("/")[-1]).exists()


def test_the_same_image_repeated_is_stored_once(tmp_path):
    """A logo on forty slides must not become forty files."""
    sink = AssetSink(tmp_path / "assets")
    first = sink.store(png_bytes("green"), ".png")
    second = sink.store(png_bytes("green"), ".png")

    assert first == second
    assert len(list((tmp_path / "assets").iterdir())) == 1


def test_figures_past_the_ceiling_are_skipped_and_counted(tmp_path, monkeypatch):
    from pipeline import extraction_assets

    monkeypatch.setattr(extraction_assets, "MAXIMUM_EXTRACTED_ASSET_BYTES", 150)
    sink = AssetSink(tmp_path / "assets")

    kept = sink.store(b"x" * 100, ".png")
    skipped = sink.store(b"y" * 100, ".png")

    assert kept and skipped == ""
    assert sink.skipped_for_size == 1


def test_figures_are_named_in_the_text_the_planner_reads():
    document = ExtractedDocument(file_name="a.docx", blocks=[
        ContentBlock(kind=BlockKind.FIGURE, text="Figure 2.1 The chain"),
    ])
    assert "[figure: Figure 2.1 The chain]" in document.as_markdown()


# --- a real run, through operations ---------------------------------------------


def extract_everything(slug):
    """Extract each source into the collection, one at a time, then confirm it --
    what someone does on the Extraction page before moving on."""
    project = operations.require_project(slug)
    for source in project.sources():
        operations.run_extraction(slug, source["name"])
    return operations.confirm_extraction(slug)


def test_a_run_stores_figures_under_the_project(project_root, tmp_path):
    slug = operations.create_project("Figures")["slug"]
    project = operations.require_project(slug)
    chapter_with_a_table_and_a_figure(project.sources_dir / "chapter.docx")
    (project.sources_dir / "figure-red.png").unlink()

    result = operations.run_extraction(slug, "chapter.docx")
    figure = next(
        block for block in result["documents"][0]["blocks"]
        if block["kind"] == BlockKind.FIGURE
    )

    assert figure["asset_path"].startswith("extraction/assets/")
    assert (project.root / figure["asset_path"]).exists()
    assert result["audit"]["figure_count"] == 1


def test_replacing_the_file_removes_figures_nothing_refers_to(project_root):
    """A picture-heavy deck replaced by another must not leave its images behind
    for good. Pictures the extraction no longer uses go -- unless what was
    transferred still points at them."""
    slug = operations.create_project("Pruning")["slug"]
    project = operations.require_project(slug)
    chapter_with_a_table_and_a_figure(project.sources_dir / "one.docx", "red")
    (project.sources_dir / "figure-red.png").unlink()
    operations.run_extraction(slug)
    assets = project.root / "extraction" / "assets"
    assert len(list(assets.iterdir())) == 1

    (project.sources_dir / "one.docx").unlink()
    chapter_with_a_table_and_a_figure(project.sources_dir / "two.docx", "blue")
    (project.sources_dir / "figure-blue.png").unlink()
    operations.run_extraction(slug)

    assert len(list(assets.iterdir())) == 1, "only the new file's picture is left"


def test_a_confirmed_figure_survives_until_the_confirmation_changes(project_root):
    """Later stages read the confirmed extraction, so its pictures stay on disk
    while the extraction moves on to a new file without them."""
    slug = operations.create_project("Keep Confirmed")["slug"]
    project = operations.require_project(slug)
    chapter_with_a_table_and_a_figure(project.sources_dir / "one.docx", "red")
    (project.sources_dir / "figure-red.png").unlink()
    extract_everything(slug)
    assets = project.root / "extraction" / "assets"

    (project.sources_dir / "one.docx").unlink()
    (project.sources_dir / "notes.md").write_text("# Notes\n\ntext", encoding="utf-8")
    operations.run_extraction(slug)
    assert len(list(assets.iterdir())) == 1, "still what the next layer reads"

    operations.confirm_extraction(slug)
    assert list(assets.iterdir()) == []


# --- one file at a time, into a collection, then confirmed -----------------------


def project_with_notes(name, *file_names):
    slug = operations.create_project(name)["slug"]
    project = operations.require_project(slug)
    for file_name in file_names:
        (project.sources_dir / file_name).write_text(
            f"# {file_name}\n\nSome text.", encoding="utf-8"
        )
    return slug, project


def names_in(extraction):
    return [document.file_name for document in extraction.documents]


def test_freshness_before_any_run(project_root):
    slug = operations.create_project("Fresh")["slug"]
    freshness = operations.extraction_freshness(operations.require_project(slug))

    assert freshness["has_run"] is False
    assert freshness["is_stale"] is False


def test_extracting_a_file_adds_only_that_file(project_root):
    from pipeline.extraction import ExtractionCollectionStore

    slug, project = project_with_notes("One At A Time", "a.md", "b.md")

    operations.run_extraction(slug, "b.md")

    assert names_in(ExtractionCollectionStore(project.root).load()) == ["b.md"]


def test_extracting_again_replaces_the_result(project_root):
    """One file, one result: a second run replaces the first, it does not add."""
    from pipeline.extraction import ExtractionCollectionStore

    slug, project = project_with_notes("Replace", "a.md")
    operations.run_extraction(slug)

    (project.sources_dir / "a.md").write_text("# a.md\n\nNew words entirely.", encoding="utf-8")
    operations.run_extraction(slug)

    collection = ExtractionCollectionStore(project.root).load()
    assert names_in(collection) == ["a.md"], "no duplicate"
    assert "New words" in collection.document("a.md").as_markdown()


def test_a_new_file_replaces_the_old_ones_result(project_root):
    """After the source is replaced, the old file's result cannot linger beside
    the new one."""
    from pipeline.extraction import ExtractionCollectionStore

    slug, project = project_with_notes("New File", "old.md")
    operations.run_extraction(slug)
    (project.sources_dir / "old.md").unlink()
    (project.sources_dir / "new.md").write_text("# new.md\n\ntext", encoding="utf-8")

    operations.run_extraction(slug)
    assert names_in(ExtractionCollectionStore(project.root).load()) == ["new.md"]

def test_later_stages_see_nothing_until_the_collection_is_confirmed(project_root):
    """The hand-off the page is built around: a half-assembled collection must not
    reach the planner as if it were the source."""
    from pipeline.extraction import ExtractionStore

    slug, project = project_with_notes("Confirm", "a.md", "b.md")
    operations.run_extraction(slug, "a.md")

    assert ExtractionStore(project.root).load().documents == []
    freshness = operations.extraction_freshness(project)
    assert freshness["pending_confirmation"] is True
    assert freshness["has_confirmed"] is False

    operations.confirm_extraction(slug)

    assert names_in(ExtractionStore(project.root).load()) == ["a.md"]
    assert operations.extraction_freshness(project)["pending_confirmation"] is False


def test_a_change_after_confirming_waits_for_the_next_confirmation(project_root):
    from pipeline.extraction import ExtractionStore

    slug, project = project_with_notes("Reconfirm", "a.md")
    extract_everything(slug)

    (project.sources_dir / "a.md").unlink()
    (project.sources_dir / "b.md").write_text("# b.md\n\ntext", encoding="utf-8")
    operations.run_extraction(slug)

    assert names_in(ExtractionStore(project.root).load()) == ["a.md"], "not until confirmed"
    assert operations.extraction_freshness(project)["pending_confirmation"] is True

    operations.confirm_extraction(slug)
    assert names_in(ExtractionStore(project.root).load()) == ["b.md"]

def test_an_empty_collection_cannot_be_confirmed(project_root):
    slug, _ = project_with_notes("Empty", "a.md")

    with pytest.raises(operations.OperationFailed, match="Extract at least one file"):
        operations.confirm_extraction(slug)


def test_a_file_that_changed_since_extraction_blocks_confirming(project_root):
    """Confirming a result for a file that has since been replaced would hand on
    something that no longer matches the file."""
    slug, project = project_with_notes("Changed", "a.md")
    operations.run_extraction(slug, "a.md")

    (project.sources_dir / "a.md").write_text("# a.md\n\nA longer, different file.",
                                              encoding="utf-8")
    freshness = operations.extraction_freshness(project)
    assert freshness["changed"] == ["a.md"]

    with pytest.raises(operations.OperationFailed, match="a.md"):
        operations.confirm_extraction(slug)


def test_a_file_not_extracted_yet_is_waiting_not_stale(project_root):
    slug, project = project_with_notes("Not Yet", "a.md")
    freshness = operations.extraction_freshness(project)

    assert freshness["added"] == ["a.md"]
    assert freshness["is_stale"] is False

def test_a_removed_source_makes_it_stale(project_root):
    slug, project = project_with_notes("Removed", "a.md", "b.md")
    extract_everything(slug)

    (project.sources_dir / "b.md").unlink()
    freshness = operations.extraction_freshness(project)

    assert freshness["removed"] == ["b.md"]
    assert freshness["is_stale"] is True


def test_only_one_file_is_extracted_at_a_time(project_root):
    """Two runs at once would each load the collection, add their file and save,
    and the second save would silently drop the first file."""
    from pipeline.extraction_run import _extraction_lock

    slug, project = project_with_notes("Busy", "a.md", "b.md")
    lock = _extraction_lock(project)
    lock.acquire()
    try:
        with pytest.raises(operations.OperationFailed, match="Another file is being extracted"):
            operations.run_extraction(slug, "b.md")
    finally:
        lock.release()

    operations.run_extraction(slug, "b.md")


def test_an_unknown_file_cannot_be_extracted(project_root):
    slug, _ = project_with_notes("Unknown", "a.md")

    with pytest.raises(operations.OperationFailed, match="no source file"):
        operations.run_extraction(slug, "../project.json")


def test_a_project_from_before_collections_keeps_what_it_extracted(project_root):
    """Only the confirmed extraction exists on disk; it becomes the collection
    rather than the page showing an empty box over content later stages use."""
    from pipeline.extraction import (
        ExtractedDocument,
        Extraction,
        ExtractionCollectionStore,
        ExtractionStore,
    )

    slug, project = project_with_notes("Older", "a.md")
    ExtractionStore(project.root).save(Extraction(
        documents=[ExtractedDocument(file_name="a.md", blocks=[
            ContentBlock(kind=BlockKind.PARAGRAPH, text="kept")])],
        extracted_at="2026-01-01T00:00:00+00:00",
    ))

    assert names_in(ExtractionCollectionStore(project.root).load()) == ["a.md"]
    assert operations.extraction_freshness(project)["pending_confirmation"] is False


def test_the_orchestrator_is_told_what_is_not_extracted(project_root):
    slug, project = project_with_notes("Told", "a.md")
    extract_everything(slug)
    (project.sources_dir / "deck.md").write_text("# Deck\n\ntext", encoding="utf-8")

    snapshot = orchestrator.pipeline_snapshot(project)

    assert "Not extracted yet: deck.md" in snapshot


def test_agents_are_told_a_file_was_not_extracted(project_root):
    """Otherwise an agent reports a deck nobody extracted as one it failed to read."""
    from pipeline.agent_runtime import stage_context
    from pipeline.stages import Stage

    slug, project = project_with_notes("Told Why", "keep.md", "deck.md")
    operations.run_extraction(slug, "keep.md")
    operations.confirm_extraction(slug)

    assert "deck.md (not extracted)" in stage_context(project, Stage.SCENE_PLAN)


# --- over HTTP --------------------------------------------------------------------


@pytest.fixture
def client(project_root):
    """The project service is a process-wide singleton holding the projects
    folder it was first built with, so each test starts it again against its own
    temporary folder -- the same reset the operations tests do."""
    from litestar.testing import TestClient

    from pipeline.web import routes_project as routes_project_module
    from pipeline.web.app import create_app

    routes_project_module._service = None
    with TestClient(app=create_app()) as test_client:
        yield test_client
    routes_project_module._service = None


def test_the_source_is_extracted_and_confirmed_over_http(project_root, client):
    slug, project = project_with_notes("Over HTTP", "a.md")

    extracted = client.post(f"/api/projects/{slug}/extraction/run", json={})
    assert extracted.status_code == 201, extracted.text
    assert extracted.json()["extracted"]["file_name"] == "a.md", "the project's one source"
    assert extracted.json()["freshness"]["added"] == []

    confirmed = client.post(f"/api/projects/{slug}/extraction/confirm")
    assert confirmed.status_code == 201, confirmed.text
    assert confirmed.json()["freshness"]["pending_confirmation"] is False
    assert client.delete(f"/api/projects/{slug}/extraction/documents/a.md").status_code \
        in (404, 405), "there is no removing one file of one"


def test_confirming_an_empty_collection_is_refused_over_http(project_root, client):
    slug, _ = project_with_notes("Refused", "a.md")

    assert client.post(f"/api/projects/{slug}/extraction/confirm").status_code == 400


def test_an_extracted_figure_is_served_inert(project_root, client):
    slug = operations.create_project("Served")["slug"]
    project = operations.require_project(slug)
    assets = project.root / "extraction" / "assets"
    assets.mkdir(parents=True)
    (assets / "abc.png").write_bytes(png_bytes())

    response = client.get(f"/api/projects/{slug}/extraction/assets/abc.png")

    assert response.status_code == 200
    assert response.headers["content-type"].startswith("image/png")
    # Bytes from someone's upload, on this app's origin: never able to run.
    assert "sandbox" in response.headers["content-security-policy"]
    assert response.headers["x-content-type-options"] == "nosniff"


def test_the_asset_route_cannot_reach_outside_the_assets_folder(project_root, client):
    slug = operations.create_project("Escape")["slug"]
    project = operations.require_project(slug)
    (project.root / "extraction" / "assets").mkdir(parents=True)

    for attempt in ("..%2F..%2Fproject.json", "..\\..\\project.json", "project.json"):
        response = client.get(f"/api/projects/{slug}/extraction/assets/{attempt}")
        assert response.status_code == 404, attempt


def test_extraction_round_trips_with_figures():
    """A content.json with no asset_path must still load."""
    old_shape = {"documents": [{"file_name": "a.docx", "blocks": [
        {"kind": "paragraph", "text": "hello", "heading_path": []},
    ]}], "extracted_at": "2026-01-01T00:00:00+00:00"}

    loaded = Extraction.model_validate(json.loads(json.dumps(old_shape)))
    assert loaded.documents[0].blocks[0].asset_path == ""
    assert loaded.options_used == {}


# --- the options do what they say ------------------------------------------------


def project_with_a_chapter(name):
    slug = operations.create_project(name)["slug"]
    project = operations.require_project(slug)
    chapter_with_a_table_and_a_figure(project.sources_dir / "chapter.docx")
    (project.sources_dir / "figure-red.png").unlink()
    return slug, project


def set_extraction_options(project, **values):
    from pipeline.stage_options import StageSettingsStore
    from pipeline.stages import Stage

    store = StageSettingsStore(project.root)
    current = store.load(Stage.EXTRACTION).values
    store.save(Stage.EXTRACTION, {**current, **values})


def test_every_option_is_read():
    """A toggle that is saved and read by nothing is indistinguishable from a
    broken one, so every declared option is one a stage reads."""
    from pipeline.stage_options import options_for
    from pipeline.stages import Stage

    assert {option.key for option in options_for(Stage.EXTRACTION)} == {
        "extract_tables", "extract_images", "extract_layout",
        "extract_backgrounds", "extract_styling", "handwriting_ocr",
    }

    # The Scene Plan's one preference is read: the language reaches both the
    # narration and the voice that speaks it. The tone is fixed: it teaches. The
    # length is typed in the plan's own setup.
    assert {option.key for option in options_for(Stage.SCENE_PLAN)} == {"language"}


def test_switching_images_off_extracts_no_figures(project_root):
    slug, project = project_with_a_chapter("No Images")
    set_extraction_options(project, extract_images=False)

    result = operations.run_extraction(slug, "chapter.docx")
    kinds = {block["kind"] for block in result["documents"][0]["blocks"]}

    assert BlockKind.FIGURE not in kinds
    assert not (project.root / "extraction" / "assets").exists()


def test_switching_images_off_later_clears_the_old_figures(project_root):
    slug, project = project_with_a_chapter("Images Later Off")
    operations.run_extraction(slug, "chapter.docx")
    assets = project.root / "extraction" / "assets"
    assert len(list(assets.iterdir())) == 1

    set_extraction_options(project, extract_images=False)
    operations.run_extraction(slug, "chapter.docx")

    assert list(assets.iterdir()) == []


def test_switching_tables_off_keeps_their_words_as_text(project_root):
    slug, project = project_with_a_chapter("No Tables")
    set_extraction_options(project, extract_tables=False)

    blocks = operations.run_extraction(slug, "chapter.docx")["documents"][0]["blocks"]

    assert all(block["kind"] != BlockKind.TABLE for block in blocks)
    flattened = next(block for block in blocks if "Wet" in block["text"])
    assert flattened["kind"] == BlockKind.PARAGRAPH
    assert flattened["text"] == "Sensor Noise Wet Low"


def test_changing_an_option_makes_that_file_stale(project_root):
    slug, project = project_with_a_chapter("Option Stale")
    extract_everything(slug)
    assert operations.extraction_freshness(project)["is_stale"] is False

    set_extraction_options(project, extract_images=False)
    freshness = operations.extraction_freshness(project)

    assert freshness["is_stale"] is True
    assert freshness["options_changed"] == ["chapter.docx"]


def test_handwriting_recognition_reads_only_pages_with_no_text(project_root, monkeypatch):
    """A page with a text layer is never sent; a textless page is read by the
    vision model and marked as transcribed."""
    from pypdf import PdfWriter

    from pipeline import llm

    slug, project = project_with_notes("Handwriting")
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    with (project.sources_dir / "scan.pdf").open("wb") as handle:
        writer.write(handle)

    sent = []
    monkeypatch.setattr("pipeline.ocr._page_images", lambda page: [b"jpeg"])
    monkeypatch.setattr(llm, "read_image_text",
                        lambda image, mime, instruction: sent.append(image) or
                        "Dear diary\n\nThe neuron fired.")
    set_extraction_options(project, handwriting_ocr=True)

    document = operations.run_extraction(slug, "scan.pdf")["documents"][0]

    assert sent == [b"jpeg"]
    assert [b["text"] for b in document["blocks"]] == ["Page 1", "Dear diary", "The neuron fired."]
    assert all(b["transcribed"] for b in document["blocks"])
    assert "1 page(s) transcribed" in document["extraction_note"]


def test_handwriting_recognition_without_a_key_still_extracts(project_root, monkeypatch):
    from pypdf import PdfWriter

    from pipeline import llm

    slug, project = project_with_notes("No Key")
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    with (project.sources_dir / "scan.pdf").open("wb") as handle:
        writer.write(handle)

    def not_configured(*_args):
        raise llm.ModelNotConfigured("No API key set")

    monkeypatch.setattr("pipeline.ocr._page_images", lambda page: [b"jpeg"])
    monkeypatch.setattr(llm, "read_image_text", not_configured)
    set_extraction_options(project, handwriting_ocr=True)

    document = operations.run_extraction(slug, "scan.pdf")["documents"][0]

    assert "handwriting recognition skipped" in document["extraction_note"]


def test_the_stage_chip_counts_documents_and_figures_not_files(project_root):
    """With figures on disk, "55 file(s)" described 54 images and a JSON file."""
    from pipeline.stages import Stage

    slug, project = project_with_a_chapter("Chip")
    extract_everything(slug)

    detail = project.stage_state(Stage.EXTRACTION).detail
    assert detail == "1 document · 1 figure"


def test_deleting_an_extracted_file_says_so_at_the_time(project_root, client):
    """Its result stays in the collection until it is removed there, and the
    delete itself says so rather than leaving it to be found on the page."""
    slug, _ = project_with_notes("Delete Extracted", "a.md", "b.md")
    operations.run_extraction(slug, "b.md")

    response = client.delete(f"/api/projects/{slug}/sources/b.md")

    assert response.status_code == 200, response.text
    assert response.json()["was_in_extraction"] is True


# --- how a deck is put together ------------------------------------------------------


def _deck_with_a_picture(tmp_path, name="design.pptx"):
    picture = tmp_path / "diagram.png"
    picture.write_bytes(png_bytes("blue"))

    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[5])
    slide.shapes.title.text = "Networks that learn"
    slide.shapes.add_picture(str(picture), Inches(5), Inches(2), width=Inches(3))
    deck.save(tmp_path / name)
    return tmp_path / name


def test_a_slides_shapes_are_captured_where_they_sit(tmp_path):
    """A picture two thirds across the slide must still be two thirds across in
    a video that is not the deck's shape, so positions are fractions."""
    from pipeline.extraction import read_pptx_design

    slides, _design = read_pptx_design(_deck_with_a_picture(tmp_path))

    shapes = slides[0].shapes
    picture = next(shape for shape in shapes if shape.kind == "picture")
    title = next(shape for shape in shapes if shape.kind == "title")

    assert slides[0].number == 1 and slides[0].title == "Networks that learn"
    assert 0 < title.width <= 1 and 0 < title.top < 1
    assert 0.4 < picture.left < 0.6, picture.left      # 5 inches across a 10-inch slide
    assert picture.z > title.z, "z-order must follow how PowerPoint stacks them"


def test_a_shape_that_declares_no_position_takes_its_layouts(tmp_path):
    """A placeholder can carry no geometry of its own. Reporting that as 0,0
    would put the title in the corner of every slide."""
    from pipeline.deck_shapes import _shape_box

    class Placeholder:
        left = top = width = height = None
        is_placeholder = True
        placeholder_format = type("Format", (), {"idx": 0})()

    class Inherited:
        left, top, width, height = 100, 200, 400, 300

    class Slide:
        slide_layout = type("Layout", (), {"placeholders": {0: Inherited()}})()

    left, top, width, height = _shape_box(Placeholder(), Slide(), 1000, 1000)

    assert (left, top, width, height) == (0.1, 0.2, 0.4, 0.3)


def test_the_decks_theme_and_proportions_are_captured(tmp_path):
    from pipeline.extraction import read_pptx_design

    _slides, design = read_pptx_design(_deck_with_a_picture(tmp_path))

    assert design.aspect_ratio == "4:3"      # python-pptx's default deck
    assert design.major_font and design.minor_font
    assert design.colours["lt1"].startswith("#") and design.colours["dk1"].startswith("#")


def test_a_system_colour_is_recorded_as_its_value_not_its_name(tmp_path):
    """`<a:sysClr val="windowText" lastClr="000000"/>` means black. Taking `val`
    recorded the theme's darkest colour as "#WINDOWTEXT"."""
    from pipeline.extraction import read_pptx_design

    _slides, design = read_pptx_design(_deck_with_a_picture(tmp_path))

    assert design.colours["dk1"] == "#000000"
    assert design.colours["lt1"] == "#FFFFFF"


def test_a_slide_that_sets_no_background_reports_the_theme_one(tmp_path):
    """Otherwise every later stage has to re-implement PowerPoint's inheritance
    to learn that a plain slide is white."""
    from pipeline.extraction import read_pptx_design

    slides, _design = read_pptx_design(_deck_with_a_picture(tmp_path))

    assert slides[0].background.colour == "#FFFFFF"
    assert slides[0].background.inherited_from == "theme"


def test_full_slide_art_on_the_layout_is_the_background(tmp_path):
    """Designed decks put their look on the layout as a full-slide picture.
    Reading only the background settings called such a deck plain white."""
    from pipeline.extraction import AssetSink, read_pptx_design

    art = tmp_path / "art.png"
    art.write_bytes(png_bytes("purple"))
    from pptx.opc.constants import RELATIONSHIP_TYPE
    from pptx.oxml.ns import qn

    deck = Presentation()
    layout = deck.slide_layouts[6]
    # python-pptx cannot add a picture to a layout, so one is made on the slide
    # and moved there, with its image related to the layout instead.
    slide = deck.slides.add_slide(layout)
    picture = slide.shapes.add_picture(
        str(art), 0, 0, width=deck.slide_width, height=deck.slide_height)
    blip = picture._element.find(".//" + qn("a:blip"))
    image_part = slide.part.related_part(blip.get(qn("r:embed")))
    blip.set(qn("r:embed"), layout.part.relate_to(image_part, RELATIONSHIP_TYPE.IMAGE))
    layout.shapes._spTree.append(picture._element)
    deck.slides.add_slide(deck.slide_layouts[5])        # a layout without art
    deck.save(tmp_path / "art.pptx")

    sink = AssetSink(tmp_path / "assets")
    slides, _design = read_pptx_design(tmp_path / "art.pptx", sink)

    assert slides[0].background.inherited_from == "layout"
    assert slides[0].background.asset_path
    assert slides[1].background.inherited_from == "theme"


def test_a_gradient_background_is_recorded_with_its_colours(tmp_path):
    from pptx.dml.color import RGBColor

    from pipeline.extraction import read_pptx_design

    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    fill = slide.background.fill
    fill.gradient()
    fill.gradient_stops[0].color.rgb = RGBColor(0x7B, 0x5C, 0xFF)
    fill.gradient_stops[1].color.rgb = RGBColor(0x2A, 0x27, 0x9C)
    deck.save(tmp_path / "gradient.pptx")

    slides, _design = read_pptx_design(tmp_path / "gradient.pptx")

    background = slides[0].background
    assert background.inherited_from == "slide"
    assert background.gradient == ["#7B5CFF", "#2A279C"]
    assert background.colour == "#7B5CFF"


def test_layout_and_design_are_captured_only_when_asked(tmp_path):
    from pipeline.extraction import extract_document

    deck_path = _deck_with_a_picture(tmp_path)

    plain = extract_document(deck_path)
    assert plain.slides == [] and plain.design is None

    with_layout = extract_document(
        deck_path, capture_layout=True,
        capture_backgrounds=True, capture_styling=True)
    assert len(with_layout.slides) == 1 and with_layout.design is not None
    assert "layout captured for 1 slide(s)" in with_layout.extraction_note


def test_capturing_a_decks_layout_does_not_change_its_text(tmp_path):
    """The words are the extraction; styling is extra and must not disturb it."""
    from pipeline.extraction import extract_document

    deck_path = _deck_with_a_picture(tmp_path)

    plain = extract_document(deck_path)
    styled = extract_document(
        deck_path, capture_layout=True,
        capture_backgrounds=True, capture_styling=True)

    assert [b.text for b in styled.blocks] == [b.text for b in plain.blocks]


def test_an_unreadable_deck_design_does_not_fail_the_extraction(tmp_path, monkeypatch):
    from pipeline import extraction

    def explode(*_args, **_kwargs):
        raise RuntimeError("theme part is corrupt")

    monkeypatch.setattr(extraction, "read_pptx_design", explode)
    document = extraction.extract_document(
        _deck_with_a_picture(tmp_path), capture_layout=True,
        capture_backgrounds=True, capture_styling=True)

    assert document.blocks, "the words must still be extracted"
    assert document.slides == []


def test_a_slide_title_is_extracted_once(tmp_path):
    """It was extracted twice -- as the slide's heading and again as a paragraph
    of its own -- because `shape is slide.shapes.title` compares two freshly
    built proxies and is never true. Every agent read every deck title twice."""
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[5])
    slide.shapes.title.text = "Networks that learn"
    slide.shapes.add_textbox(Inches(1), Inches(3), Inches(4), Inches(1)) \
        .text_frame.text = "Some body text"
    deck.save(tmp_path / "titles.pptx")

    blocks = read_pptx(tmp_path / "titles.pptx")

    assert [b.text for b in blocks] == ["Networks that learn", "Some body text"]
    assert [b.kind for b in blocks] == [BlockKind.HEADING, BlockKind.PARAGRAPH]


def test_layout_capture_does_not_copy_pictures_when_figures_are_off(tmp_path):
    """Otherwise switching images off still fills `extraction/assets/` with
    every picture in the deck, referenced by nothing a person can see."""
    from pipeline.extraction import AssetSink, extract_document

    sink = AssetSink(tmp_path / "assets")
    document = extract_document(
        _deck_with_a_picture(tmp_path), sink,
        keep_figures=False, capture_layout=True, capture_backgrounds=True, capture_styling=True)

    picture = next(shape for shape in document.slides[0].shapes if shape.kind == "picture")
    assert picture.asset_path == "", "the geometry is layout; the file is a figure"
    assert not (tmp_path / "assets").exists()


def test_a_design_only_run_does_not_claim_layout_was_captured(tmp_path):
    """Backgrounds are per slide, so a design-only run still records a slide for
    each one. Counting those as captured layout would state something untrue on
    the page right next to the switch that turned it off."""
    from pipeline.extraction import Extraction, extract_document

    document = extract_document(
        _deck_with_a_picture(tmp_path), capture_layout=False,
        capture_backgrounds=True, capture_styling=True)
    audit = Extraction(documents=[document]).audit()

    assert document.slides and document.slides[0].background is not None
    assert audit["slide_layout_count"] == 0
    assert audit["slide_background_count"] == 1


def test_backgrounds_and_styling_are_separate_switches(tmp_path):
    """A background is a picture in its own right, not a detail of the styling:
    someone who wants the slide images does not have to take the theme too."""
    from pipeline.extraction import extract_document

    deck_path = _deck_with_a_picture(tmp_path)

    backgrounds_only = extract_document(
        deck_path, capture_backgrounds=True, capture_styling=False)
    assert backgrounds_only.design is None
    assert backgrounds_only.slides[0].background.colour == "#FFFFFF", \
        "the theme still has to be read to know what is behind a plain slide"

    styling_only = extract_document(
        deck_path, capture_backgrounds=False, capture_styling=True)
    assert styling_only.design.theme_name
    assert all(slide.background is None for slide in styling_only.slides)
